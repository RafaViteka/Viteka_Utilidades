from pathlib import Path
import threading,functools,http.server,os,shutil,re,copy,json
from playwright.sync_api import sync_playwright,expect
repository=Path(__file__).resolve().parents[1]
root=repository/'frontend/dist/viteka-ui/browser'
artifacts=repository/'artifacts/ui';artifacts.mkdir(parents=True,exist_ok=True)
class Quiet(http.server.SimpleHTTPRequestHandler):
 def log_message(self,*args):pass
server=http.server.ThreadingHTTPServer(('127.0.0.1',0),functools.partial(Quiet,directory=str(root)))
threading.Thread(target=server.serve_forever,daemon=True).start()
url=f'http://127.0.0.1:{server.server_port}'
class AuthorityFixture:
 def __init__(self):
  self.accounts=[dict(id=f'{i:032x}',name=name,role=role,active=True,mustChangePassword=True,utilityIds=[],password=password) for i,(name,role,password) in enumerate([('Viteka','SuperAdministrator','jupiter'),('F. titular','Administrator','1234'),('usuario','User','0000')],1)]
  self.current=None;self.calls=[];self.prefs={};self.nix=dict(configured=False,saved=False,dsn='NIXFARMA',user='',dsnAvailable=True);self.probe_fail=False;self.schema_fail=False;self.pharmacy=None;self.tenencia_ready=False;self.stock_error=False;self.tenencia_queries=0;self.many_groups=False;self.transport_error=False;self.use_excel=False;self.sort_cases=False;self.big_group=False;self.last={};self.saved=[];self.catalog=None;self.catalog_version=0;self.catalog_custom=False
 def view(self,u):
  if u is None:return None
  ready=not u['mustChangePassword'];admin=u['role'] in ('Administrator','SuperAdministrator')
  return {**{k:v for k,v in u.items() if k!='password'},'canConfigure':ready and u['role']=='SuperAdministrator','canManageUsers':ready and admin,'canInstallUpdates':ready and admin}
 def state(self):return dict(user=self.view(self.current))
 def base_groups(self):
  groups=[dict(id='a',name='Grupo ficticio A',minimum=3,stock=1,missing=2,status='shortfall',articles=[dict(code='650001',description='Artículo ficticio A1',stock=1),dict(code='650004',description='Artículo ficticio A2',stock=0)]),dict(id='b',name='Grupo ficticio B',minimum=1,stock=3,missing=0,status='covered',requiresReview=True,selectionNote='Comprobar cada dosis.',articles=[dict(code='650002',description='Artículo con dos fichas',stock=3,issue='summed_stock')]),dict(id='c',name='Grupo ficticio C',minimum=1,stock=2,missing=0,status='covered',articles=[dict(code='650003',description='Artículo suficiente',stock=2)])]
  if self.sort_cases:groups[0].update(stock=12,missing=0,status='covered',articles=[dict(code='65000'+str(i),description=n,stock=v,issue=x) for i,n,v,x in [(1,'Z',10,None),(2,'B',0,'stock_not_found'),(3,'A',2,None),(4,'C',0,None)]])
  if self.use_excel:
   catalog=json.loads((repository/'src/Viteka.Server/Catalogs/tenencia/29.json').read_text(encoding='utf-8'))
   groups=[dict(id=g['id'],name=g['name'],minimum=g['minimum'],stock=0,missing=g['minimum'],status='shortfall',classification=g.get('atcPrefix') or g.get('classificationCode'),selectionNote=g.get('selectionNote'),sourceRow=g['sourceRow'],requiresReview=g.get('requiresReview',False),articles=[dict(code=c,description=g.get('descriptions',{}).get(c,'Artículo CN '+c),stock=0,issue='stock_not_found') for c in g['codes']]) for g in catalog['groups']]
  if self.many_groups:groups.extend(dict(id='extra-'+str(i),name='Grupo adicional '+str(i),minimum=1,stock=1,missing=0,status='covered',articles=[dict(code=str(650000+i),description='Medicamento adicional '+str(i),stock=1)]) for i in range(4,31))
  dates={'650001':('2026-12-31','2026-09-01','2026-10-20'),'650002':('2027-09-30',None,'2026-10-01')}
  for g in groups:
   for a in g['articles']:
    if a['code'] in dates:a.update(expiry=dates[a['code']][0],lastPurchase=dates[a['code']][1],lastSale=dates[a['code']][2])
  if self.big_group:
   arts=[dict(code=str(670000+i),description=f'Medicamento D{i:02d}',stock=2 if i==0 else 1 if i==1 else 0,expiry='2026-05-31' if i==0 else None,lastSale=f'2026-{(i%9)+1:02d}-15' if 2<=i<=10 else None) for i in range(13)]
   groups.append(dict(id='d',name='Grupo ficticio D',minimum=5,stock=3,missing=2,status='shortfall',articles=arts))
  return groups
 def nixfarma(self):
  nix={a['code']:(a['description'],a['stock']) for g in self.base_groups() for a in g['articles']}
  nix.update({'660001':('ADRENALINA 1 MG/ML 10 AMPOLLAS',4),'660002':('ADRENALINA 1 MG/ML 1 JERINGA',0),'660003':('NALOXONA 0,4 MG 10 AMPOLLAS',1)});return nix
 def groups(self):
  if self.catalog is None:return self.base_groups()
  nix=self.nixfarma();out=[]
  for g in self.catalog:
   arts=[dict(code=c,description=nix[c][0] if c in nix else g['descriptions'].get(c,'CN '+c),stock=nix[c][1] if c in nix else 0) for c in g['codes']]
   total=sum(a['stock'] for a in arts);out.append(dict(id=g['id'],name=g['name'],minimum=g['minimum'],stock=total,missing=max(0,g['minimum']-total),status='covered' if total>=g['minimum'] else 'shortfall',classification=g.get('classificationCode'),selectionNote=g.get('selectionNote'),articles=arts))
  return out
 def invoke(self,msg):
  cmd=msg['command'];d=msg.get('payload') or {};self.calls.append(cmd)
  def deny():raise ValueError('No tienes permiso para esta acción.')
  if cmd=='auth.state':return self.state()
  if cmd=='auth.login':
   self.current=next((u for u in self.accounts if u['name'].lower()==d['name'].lower() and u['password']==d['password'] and u['active']),None)
   if not self.current:raise ValueError('Usuario o contraseña incorrectos.')
   return self.state()
  if cmd=='auth.logout':self.current=None;return self.state()
  if cmd=='auth.password':
   if not self.current or self.current['password']!=d['currentPassword']:raise ValueError('La contraseña actual no es correcta.')
   self.current.update(password=d['newPassword'],mustChangePassword=False);return self.state()
  pref=self.prefs.setdefault(self.current['id'] if self.current else 'anonymous',dict(openFavorites=False,sidebarCollapsed=False,theme='auto',zoomPercent=100))
  if cmd=='app.state':return dict(version='0.12.0',**pref)
  if cmd=='window.appearance':return {}
  if cmd=='window.zoom':pref['zoomPercent']=d['percent'];return dict(percent=d['percent'])
  if not self.current or self.current['mustChangePassword']:deny()
  role=self.current['role'];admin=role!='User';superadmin=role=='SuperAdministrator'
  if cmd.startswith('settings.'):
   prop={'start':'openFavorites','sidebar':'sidebarCollapsed','theme':'theme'}[cmd.split('.')[1]];pref[prop]=next(iter(d.values()));return {}
  if cmd=='utilities.list':return dict(utilities=[dict(id='tenencia-obligatoria',name='Tenencia obligatoria',description='Comprueba las existencias mínimas del listado de tu provincia.')] if admin or 'tenencia-obligatoria' in self.current['utilityIds'] else [])
  if cmd=='users.list':
   if not admin:deny()
   return dict(users=[self.view(u) for u in self.accounts],utilities=[dict(id='tenencia-obligatoria',name='Tenencia obligatoria'),dict(id='fixture-a',name='Aplicación de prueba A'),dict(id='fixture-b',name='Aplicación de prueba B')])
  if cmd.startswith('users.'):
   target=next((u for u in self.accounts if u['id']==d.get('id')),None)
   if not admin or target and target['role']!='User' and not superadmin:deny()
   if cmd=='users.create':
    if d['role']!='User' and not superadmin:deny()
    self.accounts.append(dict(id=f'{len(self.accounts)+1:032x}',name=d['name'],role=d['role'],active=True,mustChangePassword=True,utilityIds=d['utilityIds'],password=d['password']));return {}
   if cmd=='users.modify':target.update(name=d['name'],role=d['role'],active=d['active'],utilityIds=d['utilityIds']);return self.state()
   if cmd=='users.password':target.update(password=d['password'],mustChangePassword=True);return {}
  if cmd.startswith('pharmacy.'):
   if not admin:deny()
   if cmd=='pharmacy.save':self.pharmacy=copy.deepcopy(d)
   return dict(profile=copy.deepcopy(self.pharmacy),provinces=[dict(code='29',name='Málaga'),dict(code='28',name='Madrid')])
  if cmd in ('file.save','file.pdf'):
   assert re.fullmatch(r'[\w\-. ]{1,80}\.(csv|xlsx|pdf)',d['name']),d['name'];self.saved.append(copy.deepcopy(d));return dict(saved=True,file=d['name'])
  if cmd.startswith('tenencia.catalog.'):
   if not superadmin:deny()
   action=cmd[len('tenencia.catalog.'):]
   if self.catalog is None:self.catalog=[dict(id=g['id'],name=g['name'],minimum=g['minimum'],codes=[a['code'] for a in g['articles']],descriptions={a['code']:a['description'] for a in g['articles']},classificationCode=g.get('classification'),selectionNote=g.get('selectionNote')) for g in self.base_groups()]
   find=lambda i:next(g for g in self.catalog if g['id']==i)
   if action=='search':
    nix=self.nixfarma()
    if d['mode']=='code':codes=list(dict.fromkeys(re.findall(r'\d{6,7}',d['text'])));arts=[dict(code=c,description=nix[c][0],stock=nix[c][1]) if c in nix else dict(code=c,description='',stock=None,issue='not_found') for c in codes]
    elif d['mode']=='description':words=d['text'].upper().split();arts=[dict(code=c,description=v[0],stock=v[1]) for c,v in nix.items() if all(w in v[0].upper() for w in words)]
    else:
     assert d['field']=='AGRUP_HOMOGENEA',d;arts=[dict(code=c,description=v[0],stock=v[1]) for c,v in nix.items() if d['text'].upper()=='AH0001' and c in ('660001','660002')]
    return dict(code='ok',message=f'{len(arts)} artículos para la búsqueda.',articles=arts,truncated=False)
   if action=='group.save':
    assert d['name'].strip() and 1<=d['minimum']<=10000,d
    if d.get('id'):find(d['id']).update(name=d['name'],minimum=d['minimum'],classificationCode=d['code'] or None,selectionNote=d['note'] or None)
    else:
     g=dict(id='custom-'+str(len(self.catalog)+1),name=d['name'],minimum=d['minimum'],codes=[],descriptions={},classificationCode=d['code'] or None,selectionNote=d['note'] or None)
     at=next((i for i,x in enumerate(self.catalog) if x['id']==d.get('afterId')),len(self.catalog)-1);self.catalog.insert(at+1,g)
   elif action=='group.delete':self.catalog.remove(find(d['id']))
   elif action=='codes.add':
    assert 1<=len(d['items'])<=50,len(d['items']);g=find(d['id'])
    for it in d['items']:
     assert re.fullmatch(r'\d{6,7}',it['code'])
     if it['code'] not in g['codes']:g['codes'].append(it['code'])
     if it['description']:g['descriptions'][it['code']]=it['description']
   elif action=='codes.remove':g=find(d['id']);g['codes']=[c for c in g['codes'] if c not in d['codes']]
   elif action=='reset':self.catalog=None;self.catalog_custom=False;self.catalog_version+=1;self.last={};return dict(list=dict(version='Fixture de pruebas',source='Fixture',groups=[dict(id=g['id'],name=g['name'],minimum=g['minimum'],codes=[a['code'] for a in g['articles']],descriptions={a['code']:a['description'] for a in g['articles']}) for g in self.base_groups()]),customized=False,provinceName='Málaga')
   if action!='state':self.catalog_custom=True;self.catalog_version+=1
   return dict(list=dict(version='Personalizada-'+str(self.catalog_version) if self.catalog_custom else 'Fixture de pruebas',source='Fixture',groups=copy.deepcopy(self.catalog)),customized=self.catalog_custom,provinceName='Málaga')
  if cmd.startswith('tenencia.'):
   if cmd=='tenencia.schema':
    if not superadmin:deny()
    return dict(code='column',message='Comprobación de tablas: Oracle no reconoce una columna. Diagnóstico: column · SQLSTATE 42000 · ORA-00904') if self.schema_fail else dict(code='schema_verified',message='Estructura verificada: ARTICULOS y EXISTENCIAS; columna de stock ACTUAL.')
   if not admin and 'tenencia-obligatoria' not in self.current['utilityIds']:deny()
   if self.pharmacy is None:return dict(code='pharmacy_pending',message='El titular o Viteka deben guardar la provincia en Datos de la farmacia.',province='')
   if not self.tenencia_ready:return dict(code='catalog_pending',message='Lista de tenencia pendiente para Málaga. No se ha evaluado el cumplimiento.',province='29',provinceName='Málaga')
   groups=self.groups()
   preview=[dict(**{**r,'stock':None,'missing':None,'status':'not_checked','articles':[dict(**{**a,'stock':None,'issue':'not_checked'}) for a in r['articles']]}) for r in groups]
   key=(self.many_groups,self.use_excel,self.sort_cases,self.big_group,self.catalog_version)
   response=dict(code='ready',message='Lista disponible. Pulsa Consultar para leer las existencias actuales.',province=self.pharmacy['province'],provinceName='Málaga',listVersion='Personalizada-'+str(self.catalog_version) if self.catalog_custom else 'Fixture de pruebas',source='Datos ficticios exclusivos del smoke UI',rows=preview,fields=dict(expiry=True,movements=True,homogeneous=['AGRUP_HOMOGENEA']),customized=self.catalog_custom)
   if cmd=='tenencia.state' and key in self.last:return {**copy.deepcopy(self.last[key]),'message':'Última consulta guardada. Pulsa Consultar existencias para actualizarla.','cached':True}
   if cmd=='tenencia.query':
    self.tenencia_queries+=1
    if self.transport_error:raise ValueError('Servicio no disponible.')
    if self.stock_error:return dict(**{**response,'code':'network','message':'No se alcanza el servicio Oracle.'})
    response.update(code='ok',message='Consulta de solo lectura completada.',checkedAt='2026-10-28T12:00:00Z',rows=groups);self.last[key]=copy.deepcopy(response)
   return response
  if cmd.startswith('nixfarma.'):
   if not superadmin:deny()
   assert not d,'Oracle credentials/SQL must not travel through the web bridge'
   if cmd=='nixfarma.status':return copy.deepcopy(self.nix)
   if cmd=='nixfarma.configure':self.nix.update(configured=True,saved=True,user='fixture-oracle');return dict(state=copy.deepcopy(self.nix))
   if cmd=='nixfarma.forget':self.nix.update(configured=False,saved=False,user='');return dict(state=copy.deepcopy(self.nix))
   if cmd=='nixfarma.test':return dict(success=not self.probe_fail,message='Oracle ha rechazado el usuario o la contraseña.' if self.probe_fail else 'Conexión verificada: prueba de solo lectura correcta.',sqlState='28000' if self.probe_fail else None,oracleCode=1017 if self.probe_fail else None)
  if cmd.startswith('update.'):
   if cmd in ('update.state','update.check'):return dict(canInstallHere=True,available=True,version='0.13.0',lastChecked='2026-10-28T10:00:00Z')
   if cmd=='update.install':
    if not admin:deny()
    return dict(cancelled=True)
   if cmd=='update.publication':return dict(opened=True)
   if cmd=='update.history':return dict(local=[dict(kind='version_started',version='0.12.0',at='2026-10-28T10:00:00Z')],published=[dict(title='Viteka 0.12.0',version='0.12.0',notes='<script>window.historyScriptRan=true</script>',publishedAt='2026-10-28T10:00:00Z')])
  if cmd=='resource.open':return dict(opened=d['id']!='anydesk')
  if cmd=='clipboard.copy':return {}
  raise ValueError('Acción no disponible.')
fixture=AuthorityFixture()
with sync_playwright() as p:
 executable=os.environ.get('VITEKA_CHROME') or shutil.which('chromium') or shutil.which('google-chrome') or str(Path(os.environ.get('PROGRAMFILES','C:/Program Files'))/'Google/Chrome/Application/chrome.exe')
 browser=p.chromium.launch(executable_path=executable,headless=True)
 context=browser.new_context(viewport={'width':1240,'height':820})
 def native_call(source,msg):
  try:return dict(id=msg['id'],ok=True,result=fixture.invoke(msg))
  except ValueError as e:return dict(id=msg['id'],ok=False,error=str(e))
 context.expose_binding('fixtureNative',native_call)
 context.add_init_script('''window.chrome=window.chrome||{};let listener;window.chrome.webview={addEventListener:(type,fn)=>listener=fn,postMessage:msg=>{window.fixtureNative(msg).then(response=>{if(msg.command==='window.zoom'&&response.ok)document.documentElement.style.zoom=String(response.result.percent/100);listener({data:response});});}};''')
 page=context.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)));page.on('console',lambda e:errors.append(e.text) if e.type=='error' else None)
 page.route('**/favicon.ico',lambda route:route.fulfill(status=204));page.goto(url)
 page.get_by_role('heading',name='Iniciar sesión',exact=True).wait_for()
 assert page.get_by_role('navigation').count()==0
 assert page.get_by_role('heading',name='Conexión a NixFarma').count()==0
 page.screenshot(path=str(artifacts/'angular-login.png'),full_page=True)
 def login(name,password,new_password=None):
  page.get_by_label('Usuario',exact=True).fill(name);page.get_by_label('Contraseña',exact=True).fill(password);page.get_by_role('button',name='Entrar',exact=True).click()
  if new_password:
   page.get_by_role('heading',name='Cambia tu contraseña',exact=True).wait_for()
   assert page.get_by_role('navigation').count()==0
   page.get_by_label('Contraseña actual',exact=True).fill(password);page.get_by_label('Nueva contraseña',exact=True).fill(new_password);page.get_by_label('Repetir nueva contraseña',exact=True).fill(new_password);page.get_by_role('button',name='Guardar contraseña',exact=True).click()
  page.get_by_role('heading',name='Portal de soporte',exact=True).wait_for()
  page.locator('.layout[data-ready=true]').wait_for();page.wait_for_timeout(150)
  prompt=page.get_by_role('dialog',name='Actualización disponible',exact=True)
  shown=prompt.count()>0
  if shown:
   prompt.get_by_text('¿Quieres actualizar ahora?',exact=False).wait_for();expect(prompt.get_by_role('button',name='Actualizar ahora',exact=True)).to_be_visible()
   if not (artifacts/'angular-update-prompt.png').exists() or name=='Viteka':page.screenshot(path=str(artifacts/'angular-update-prompt.png'),full_page=True)
   prompt.get_by_role('button',name='Más tarde',exact=True).click();prompt.wait_for(state='hidden')
  return shown
 def go(name):
  if name in ('Mi cuenta','Usuarios','Actualizaciones','Auditoría'):
   go('Configuración');page.locator('.setting-row').filter(has=page.get_by_text(name,exact=True)).click()
  else:page.get_by_role('navigation').get_by_role('button',name=name,exact=True).click()
  page.get_by_role('heading',name=name,exact=True).wait_for()
 def appearance(theme):
  page.locator('.setting-row').filter(has=page.get_by_text('Apariencia',exact=True)).click();dialog=page.get_by_role('dialog',name='Apariencia',exact=True);dialog.get_by_label('Apariencia',exact=True).select_option(theme);dialog.get_by_role('button',name='Guardar cambios',exact=True).click();dialog.wait_for(state='hidden')
 assert login('Viteka','jupiter','FixtureSuperPass1'),'falta la ventana de actualización'
 assert page.get_by_label('Apariencia',exact=True).count()==0
 logout=page.get_by_role('button',name='Cerrar sesión',exact=True);brand=page.get_by_role('button',name='Web Viteka',exact=True);assert logout.bounding_box()['y']>page.get_by_role('navigation').bounding_box()['y']+page.get_by_role('navigation').bounding_box()['height'];assert logout.bounding_box()['y']+logout.bounding_box()['height']+12<=brand.bounding_box()['y'];assert page.locator('.logout-button v-icon svg').count()==1
 assert page.get_by_role('navigation').get_by_role('button',name='Actualizaciones',exact=True).count()==0
 page.screenshot(path=str(artifacts/'angular-desktop.png'),full_page=True)
 page.keyboard.down('Control');page.mouse.wheel(0,-120);page.keyboard.up('Control');page.wait_for_function("() => document.documentElement.style.zoom==='1.1'")
 page.keyboard.press('Control+0');page.wait_for_function("() => document.documentElement.style.zoom==='1'")
 page.get_by_role('button',name='Cerrar menú lateral',exact=True).click();page.wait_for_function("() => Math.round(document.querySelector('aside').getBoundingClientRect().width)===76")
 page.reload();page.get_by_role('button',name='Abrir menú lateral',exact=True).wait_for();page.locator('.layout[data-ready=true]').wait_for();page.wait_for_timeout(150);assert page.get_by_role('dialog',name='Actualización disponible',exact=True).count()==0;page.get_by_role('button',name='Abrir menú lateral',exact=True).click()
 go('Mi cuenta');appearance('dark');page.wait_for_function("() => document.documentElement.dataset.theme==='dark'")
 page.reload();page.wait_for_function("() => document.documentElement.dataset.theme==='dark'")
 page.get_by_role('navigation').get_by_role('button',name='Inicio',exact=True).click();page.get_by_role('button',name='Cerrar menú lateral',exact=True).click();page.wait_for_function("() => Math.round(document.querySelector('aside').getBoundingClientRect().width)===76")
 assert page.locator('.logout-button').bounding_box()['width']<=60
 assert page.locator('.brand-bottom .isologo').evaluate('el=>getComputedStyle(el).filter')=='brightness(0) invert(1)'
 page.screenshot(path=str(artifacts/'angular-dark.png'),full_page=True)
 page.get_by_role('button',name='Abrir menú lateral',exact=True).click();go('Configuración')
 assert page.get_by_role('button',name='Buscar actualizaciones',exact=True).count()==0
 assert page.get_by_label('Apariencia',exact=True).count()==0
 assert page.locator('input[type=password]').count()==0
 page.locator('.setting-row').filter(has=page.get_by_text('Conexión a NixFarma',exact=True)).click()
 page.get_by_role('button',name='Configurar conexión',exact=True).click();page.get_by_text('Conexión guardada. Pulsa Probar conexión para verificarla.',exact=True).wait_for()
 page.get_by_role('button',name='Probar conexión',exact=True).click();page.get_by_text('Conexión verificada: prueba de solo lectura correcta.',exact=True).wait_for()
 fixture.schema_fail=True;page.get_by_role('button',name='Comprobar tablas de tenencia',exact=True).click();page.get_by_text('Comprobación de tablas: Oracle no reconoce una columna. Diagnóstico: column · SQLSTATE 42000 · ORA-00904',exact=True).wait_for();expect(page.locator('.connection-result.verified')).to_have_count(0);assert page.get_by_text('Conexión verificada: prueba de solo lectura correcta.',exact=True).count()==0
 page.get_by_role('button',name='Probar conexión',exact=True).click();page.get_by_text('Conexión verificada: prueba de solo lectura correcta.',exact=True).wait_for();assert page.get_by_text('Comprobación de tablas: Oracle no reconoce una columna. Diagnóstico: column · SQLSTATE 42000 · ORA-00904',exact=True).count()==0
 fixture.probe_fail=True;page.get_by_role('button',name='Probar conexión',exact=True).click();page.get_by_text('Diagnóstico: 28000 · ORA-01017',exact=True).wait_for();expect(page.locator('.connection-result.verified')).to_have_count(0)
 page.keyboard.press('Escape');go('Auditoría');page.get_by_role('heading',name='Todavía no hay registros de auditoría',exact=True).wait_for();assert page.get_by_role('navigation').get_by_role('button',name='Auditoría',exact=True).count()==0;go('Configuración');page.screenshot(path=str(artifacts/'angular-settings.png'),full_page=True);go('Actualizaciones');page.get_by_role('button',name='Buscar actualizaciones',exact=True).click();page.get_by_text('Versión 0.13.0 disponible.',exact=True).wait_for()
 assert page.get_by_role('navigation').get_by_role('button',name='Configuración',exact=True).get_attribute('aria-current')=='page'
 page.get_by_role('button',name='Actualizar ahora',exact=True).click();page.get_by_text('Instalación cancelada. Puedes seguir usando esta versión.',exact=True).wait_for()
 page.get_by_role('button',name='Historial de actualizaciones',exact=True).click();page.get_by_role('dialog',name='Historial de actualizaciones',exact=True).wait_for();page.get_by_text('Versión ejecutada 0.12.0',exact=True).wait_for()
 page.get_by_role('dialog').get_by_role('group').count() if False else None
 page.get_by_role('dialog').locator('summary').click();assert '<script>' in page.get_by_role('dialog').locator('pre').inner_text();assert page.evaluate('window.historyScriptRan===undefined')
 page.keyboard.press('Escape');page.get_by_role('dialog').wait_for(state='hidden')
 go('Aplicaciones');page.locator('button.resource',has_text='Tenencia obligatoria').click();page.get_by_role('heading',name='Tenencia obligatoria',exact=True).wait_for();page.get_by_text('El titular o Viteka deben guardar la provincia en Datos de la farmacia.',exact=True).wait_for();assert fixture.tenencia_queries==0
 go('Configuración');page.locator('.setting-row').filter(has=page.get_by_text('Datos de la farmacia',exact=True)).click();pharmacy=page.get_by_role('dialog',name='Datos de la farmacia',exact=True);pharmacy.get_by_label('Nombre de la farmacia',exact=True).fill('Farmacia ficticia');pharmacy.get_by_label('Provincia',exact=True).select_option('29');pharmacy.get_by_label('Localidad',exact=True).fill('Málaga');pharmacy.get_by_label('Correo electrónico',exact=True).fill('fixture@example.test');pharmacy.get_by_role('button',name='Guardar datos',exact=True).click();pharmacy.get_by_text('Datos guardados para toda la farmacia. La próxima consulta utilizará esta provincia.',exact=True).wait_for();page.screenshot(path=str(artifacts/'angular-pharmacy.png'),full_page=True);page.keyboard.press('Escape');pharmacy.wait_for(state='hidden')
 go('Aplicaciones');page.locator('button.resource',has_text='Tenencia obligatoria').click();page.get_by_text('Lista de tenencia pendiente para Málaga. No se ha evaluado el cumplimiento.',exact=True).wait_for();expect(page.get_by_role('button',name='Consultar existencias',exact=True)).to_be_disabled();assert page.locator('table').count()==0
 fixture.tenencia_ready=True;page.get_by_role('button',name='Revisar disponibilidad',exact=True).click();expect(page.get_by_role('button',name='Consultar existencias',exact=True)).to_be_enabled();assert fixture.tenencia_queries==0
 expect(page.locator('.group-button')).to_have_count(3);expect(page.locator('.tenencia-medications tbody tr')).to_have_count(2);expect(page.locator('.drug-stock').first).to_have_text('Sin consultar');page.get_by_role('button',name='Seleccionar grupo Grupo ficticio B',exact=True).click();expect(page.locator('.tenencia-medications tbody tr')).to_have_count(1);page.get_by_role('button',name='Seleccionar grupo Grupo ficticio A',exact=True).click()
 page.get_by_role('button',name='Consultar existencias',exact=True).click();page.get_by_role('button',name='Ver detalle de Grupo ficticio A',exact=True).wait_for();assert fixture.tenencia_queries==1;page.screenshot(path=str(artifacts/'angular-tenencia.png'),full_page=True)
 fixture.many_groups=True;page.get_by_role('button',name='Revisar disponibilidad',exact=True).click();expect(page.locator('.group-button')).to_have_count(30);page.get_by_role('button',name='Seleccionar grupo Grupo adicional 30',exact=True).click();page.get_by_role('region',name='Medicamentos del grupo',exact=True).get_by_text('Medicamento adicional 30',exact=True).wait_for();expect(page.locator('.drug-stock').first).to_have_text('Sin consultar');fixture.many_groups=False;page.get_by_role('button',name='Revisar disponibilidad',exact=True).click();expect(page.locator('.group-button')).to_have_count(3);page.get_by_role('button',name='Consultar existencias',exact=True).click();page.get_by_role('button',name='Ver detalle de Grupo ficticio A',exact=True).wait_for()
 page.set_viewport_size({'width':820,'height':680});assert page.evaluate('document.documentElement.scrollWidth<=innerWidth');page.set_viewport_size({'width':1240,'height':820})
 page.get_by_role('button',name='Seleccionar grupo Grupo ficticio B',exact=True).click();page.get_by_role('button',name='Ver detalle de Grupo ficticio B',exact=True).click();page.get_by_role('dialog',name='Detalle de tenencia',exact=True).get_by_text('Varias fichas de existencias: se muestra la suma.',exact=True).wait_for();page.keyboard.press('Escape');expect(page.get_by_role('dialog')).to_have_count(0)
 page.get_by_label('Estado',exact=True).select_option('shortfall');expect(page.locator('.group-button')).to_have_count(1);expect(page.locator('.tenencia-medications tbody tr')).to_have_count(2);page.get_by_label('Buscar grupo, artículo o CN',exact=True).fill('650099');page.get_by_text('No hay grupos que coincidan con estos filtros.',exact=True).wait_for();page.get_by_label('Buscar grupo, artículo o CN',exact=True).fill('');page.get_by_label('Estado',exact=True).select_option('all')
 fixture.stock_error=True;page.get_by_role('button',name='Consultar existencias',exact=True).click();page.get_by_text('No se alcanza el servicio Oracle.',exact=True).wait_for();expect(page.locator('.group-button')).to_have_count(3);expect(page.locator('.drug-stock').first).to_have_text('Sin consultar');fixture.stock_error=False
 page.get_by_label('Estado',exact=True).select_option('covered');fixture.transport_error=True;page.get_by_role('button',name='Consultar existencias',exact=True).click();page.get_by_text('Servicio no disponible.',exact=True).wait_for();expect(page.locator('.group-button')).to_have_count(3);expect(page.locator('.drug-stock').first).to_have_text('Sin consultar');fixture.transport_error=False
 fixture.sort_cases=True;go('Aplicaciones');page.locator('button.resource',has_text='Tenencia obligatoria').click();page.get_by_role('button',name='Consultar existencias',exact=True).click();expect(page.get_by_role('button',name='Seleccionar grupo Grupo ficticio B',exact=True)).to_have_attribute('aria-pressed','true');page.get_by_role('button',name='Seleccionar grupo Grupo ficticio A',exact=True).click();expect(page.locator('.drug-stock')).to_have_count(4);expect(page.locator('.drug-stock')).to_have_text(['10','2','0','0']);assert all(t.strip()=='' for t in page.locator('.groups-table .row-state').all_inner_texts());assert page.locator('.groups-table .row-state').first.get_attribute('aria-label') in ('Correcto','Falta');assert page.get_by_label('Estado',exact=True).locator('option').all_inner_texts()==['Todos','Sin consultar','Falta','Correcto']
 meds=page.get_by_role('region',name='Medicamentos del grupo',exact=True)
 sortstock=meds.get_by_role('columnheader',name=re.compile('Existencias')).get_by_role('button');sortstock.click();expect(page.locator('.drug-stock')).to_have_text(['0','0','2','10']);sortstock.click();expect(page.locator('.drug-stock')).to_have_text(['10','2','0','0']);assert page.locator('.drug-stock',has_text='Pendiente').count()==0
 meds.get_by_role('columnheader',name=re.compile('Medicamento')).get_by_role('button').click();expect(page.locator('.article-select')).to_have_text(['A','B','C','Z'])
 page.get_by_text('Sin ficha de existencias en NixFarma: cuenta como 0.',exact=True).wait_for();assert page.get_by_text('Pendiente no significa cero',exact=False).count()==0
 groups=page.get_by_role('region',name='Grupos de tenencia',exact=True);groups.get_by_role('columnheader',name=re.compile('Mín.')).get_by_role('button').click();expect(page.locator('.groups-table tbody tr td:nth-child(3)')).to_have_text(['1','1','3']);assert page.get_by_role('button',name='Seleccionar grupo Grupo ficticio A',exact=True).get_attribute('aria-pressed')=='true';fixture.sort_cases=False
 page.get_by_role('button',name='Consultar existencias',exact=True).click();page.get_by_role('button',name='Ver detalle de Grupo ficticio A',exact=True).wait_for();expect(page.locator('.tenencia-summary')).to_contain_text('2 de 3 correctos')
 chip=page.locator('.tenencia-summary button.summary-chip').nth(1);chip.click();expect(page.get_by_label('Estado',exact=True)).to_have_value('shortfall');chip.click();expect(page.get_by_label('Estado',exact=True)).to_have_value('all')
 page.get_by_role('button',name='Seleccionar grupo Grupo ficticio C',exact=True).click();page.locator('.tenencia-information').click();page.keyboard.press('f');expect(page.locator('.group-button[aria-pressed=true]')).to_have_count(1);expect(page.locator('.groups-table tr.chosen .row-state')).to_have_attribute('aria-label','Falta')
 # Última consulta memorizada con fecha y hora; columnas Mín./Estado centradas; caducidad solo con existencias.
 page.get_by_text(re.compile(r'Última consulta: 28/10/2026 \d\d:\d\d')).wait_for()
 page.get_by_role('button',name='Revisar disponibilidad',exact=True).click();page.get_by_text('Última consulta guardada. Pulsa Consultar existencias para actualizarla.',exact=True).wait_for();expect(page.locator('.checked-at.cached')).to_have_count(1);expect(page.locator('.groups-table .row-state')).to_have_count(3);assert fixture.tenencia_queries>=1
 for sel in ['.groups-table thead th:nth-child(3)','.groups-table thead th:nth-child(4)','.groups-table tbody td:nth-child(3)','.groups-table tbody td:nth-child(4)']:assert page.locator(sel).first.evaluate('e=>getComputedStyle(e).textAlign')=='center',sel
 page.get_by_role('button',name='Seleccionar grupo Grupo ficticio A',exact=True).click();meds=page.get_by_role('region',name='Medicamentos del grupo',exact=True)
 expect(meds.locator('.expiry')).to_have_count(1);expect(meds.locator('.expiry').first).to_have_attribute('data-expiry','soon');expect(meds.locator('.expiry').first).to_have_text('12/2026')
 page.get_by_role('button',name='Seleccionar grupo Grupo ficticio B',exact=True).click();expect(meds.locator('.expiry').first).to_have_attribute('data-expiry','ok');page.get_by_role('button',name='Seleccionar grupo Grupo ficticio C',exact=True).click();expect(meds.locator('.expiry').first).to_have_attribute('data-expiry','none');expect(meds.locator('.expiry').first).to_have_text('Sin fecha')
 page.get_by_role('button',name='Seleccionar grupo Grupo ficticio A',exact=True).click();page.get_by_role('button',name='Seleccionar artículo 650001',exact=True).click();page.get_by_role('region',name='Información de la selección',exact=True).get_by_text('Compra 01/09/2026 · Venta 20/10/2026',exact=True).wait_for()
 page.screenshot(path=str(artifacts/'angular-tenencia-expiry.png'),full_page=True)
 # Exportaciones: Excel con logo, CSV y PDF a través del puente nativo.
 import base64,zipfile,io
 def export(kind):
  page.locator('details.export-menu summary').click();page.get_by_role('menuitem',name=re.compile('^'+kind)).click();page.wait_for_function('n=>true',arg=0);page.wait_for_timeout(300);return fixture.saved[-1]
 page.locator('details.export-menu summary').click();page.screenshot(path=str(artifacts/'angular-tenencia-export.png'),full_page=True);page.keyboard.press('Escape');page.locator('details.export-menu summary').click()
 x=export('Excel');assert x['name'].startswith('tenencia-faltas-') and x['name'].endswith('.xlsx'),x['name']
 book=zipfile.ZipFile(io.BytesIO(base64.b64decode(x['base64'])));names=book.namelist();assert 'xl/media/image1.png' in names and 'xl/worksheets/sheet2.xml' in names,names;sheet=book.read('xl/worksheets/sheet1.xml').decode();assert 'Unidades que faltan'.lower() in sheet.lower() and 'Grupo ficticio A' in sheet and '<v>2</v>' in sheet,sheet[:400];assert 'Grupo ficticio B' not in sheet
 try:
  import openpyxl;wb=openpyxl.load_workbook(io.BytesIO(base64.b64decode(x['base64'])));assert wb.sheetnames==['Faltas','Todos los grupos'],wb.sheetnames;assert wb['Faltas']['A8'].value=='Código' and wb['Faltas']['E9'].value==2
 except ImportError:pass
 c=export('CSV');csv=base64.b64decode(c['base64']).decode('utf-8');assert csv.startswith('\ufeffCódigo;Grupo;Mínimo;Existencias;Faltan;CN;Medicamento;Existencias CN;Caducidad'),csv[:120];assert c['name'].endswith('.csv') and '650001' in csv and '12/2026' in csv
 f=export('PDF');assert f['name'].startswith('tenencia-faltas-') and f['name'].endswith('.pdf') and 'base64' not in f
 page.emulate_media(media='print');page.screenshot(path=str(artifacts/'angular-tenencia-print.png'),full_page=True);assert page.locator('.tenencia-print img[src="logotipo.png"]').count()==1;expect(page.locator('.tenencia-print .group-line .missing').first).to_have_text('2');page.emulate_media(media='screen')
 # Grupos con más de 10 artículos: los de existencias y las 8 compras/ventas más recientes; «Ver todos».
 fixture.big_group=True;page.get_by_role('button',name='Revisar disponibilidad',exact=True).click();page.get_by_role('button',name='Consultar existencias',exact=True).click();page.get_by_role('button',name='Seleccionar grupo Grupo ficticio D',exact=True).click()
 expect(meds.locator('tbody tr')).to_have_count(10);assert 'Medicamento D00' in meds.inner_text() and 'Medicamento D11' not in meds.inner_text() and 'Medicamento D09' not in meds.inner_text() and 'Medicamento D02' in meds.inner_text();expect(meds.locator('.expiry[data-expiry=expired]')).to_have_count(1)
 meds.get_by_role('button',name='Ver todos (13)',exact=True).click();expect(meds.locator('tbody tr')).to_have_count(13);meds.get_by_role('button',name='Ver solo los más recientes',exact=True).click();expect(meds.locator('tbody tr')).to_have_count(10);page.screenshot(path=str(artifacts/'angular-tenencia-top10.png'),full_page=True);fixture.big_group=False
 # Edición del listado (solo superadministrador).
 page.get_by_role('button',name='Editar listado',exact=False).click();page.get_by_role('region',name='Edición del listado',exact=True).wait_for();expect(page.locator('.editor-groups tbody tr')).to_have_count(3)
 page.get_by_role('button',name='Añadir grupo',exact=True).click();dlg=page.get_by_role('dialog',name='Nuevo grupo',exact=True);dlg.get_by_label('Código',exact=True).fill('X01');dlg.get_by_label('Grupo',exact=True).fill('Grupo nuevo de prueba');dlg.get_by_label('Mínimo',exact=True).fill('2');dlg.get_by_role('button',name='Guardar',exact=True).click();dlg.wait_for(state='hidden')
 expect(page.locator('.editor-groups tbody tr')).to_have_count(4);expect(page.locator('.editor-groups tr.chosen')).to_contain_text('Grupo nuevo de prueba')
 ed=page.get_by_role('region',name='Artículos del grupo',exact=True);ed.get_by_role('button',name='Por descripción',exact=True).click();ed.get_by_label('Descripción',exact=True).fill('adrenalina');ed.get_by_role('button',name='Buscar en NixFarma',exact=True).click();expect(ed.locator('.results tbody tr')).to_have_count(2);page.screenshot(path=str(artifacts/'angular-tenencia-editor.png'),full_page=True)
 ed.get_by_role('button',name='Añadir al grupo',exact=True).click();expect(ed.locator('.articles tbody tr')).to_have_count(2);assert fixture.catalog[1]['codes']==['660001','660002'] or any(g['codes']==['660001','660002'] for g in fixture.catalog)
 ed.get_by_role('button',name='Por agrupación homogénea',exact=True).click();ed.get_by_label('Código de agrupación homogénea',exact=True).fill('AH0001');ed.get_by_role('button',name='Buscar en NixFarma',exact=True).click();ed.get_by_text('Ya en el grupo',exact=True).first.wait_for()
 ed.get_by_role('button',name='Por código nacional',exact=True).click();ed.get_by_label('Códigos nacionales',exact=True).fill('660003, 699999');ed.get_by_role('button',name='Buscar en NixFarma',exact=True).click();ed.get_by_text('No dado de alta en NixFarma',exact=True).wait_for();ed.get_by_role('button',name='Añadir al grupo',exact=True).click();expect(ed.locator('.articles tbody tr')).to_have_count(3)
 ed.get_by_role('button',name='Quitar 660002',exact=True).click();expect(ed.locator('.articles tbody tr')).to_have_count(2)
 page.get_by_role('button',name='Editar grupo Grupo nuevo de prueba',exact=True).click();dlg=page.get_by_role('dialog',name='Editar grupo',exact=True);dlg.get_by_label('Mínimo',exact=True).fill('6');dlg.get_by_label('Nota',exact=True).fill('Nota de prueba');dlg.get_by_role('button',name='Guardar',exact=True).click();dlg.wait_for(state='hidden');ed.get_by_text(re.compile('Mínimo 6')).wait_for()
 page.get_by_role('button',name='Importar códigos',exact=True).click();imp=page.get_by_role('dialog',name='Importar códigos',exact=True);imp.get_by_label('Filas a importar',exact=True).fill('X01;650003\nZ99;Grupo importado;3;650004;ARTÍCULO IMPORTADO\nlinea mala');imp.get_by_role('button',name='Importar',exact=True).click();imp.get_by_text('2 códigos añadidos en 2 grupos · 1 grupos nuevos · 1 líneas no reconocidas.',exact=True).wait_for();imp.get_by_role('button',name='Cerrar',exact=True).last.click()
 expect(page.locator('.editor-groups tbody tr')).to_have_count(5)
 page.get_by_role('button',name='Exportar Excel',exact=True).click();page.wait_for_timeout(300);x=fixture.saved[-1];assert x['name'].startswith('tenencia-codigos-') and x['name'].endswith('.xlsx');book=zipfile.ZipFile(io.BytesIO(base64.b64decode(x['base64'])));assert 'Grupo importado' in book.read('xl/worksheets/sheet1.xml').decode() and 'xl/media/image1.png' in book.namelist()
 page.get_by_role('button',name='CSV',exact=True).click();page.wait_for_timeout(300);csv=base64.b64decode(fixture.saved[-1]['base64']).decode('utf-8');assert csv.startswith('\ufeffCódigo grupo;Grupo;Mínimo;CN;Descripción;Nota') and 'X01;Grupo nuevo de prueba;6;660001' in csv,csv[:300]
 page.get_by_role('button',name='Eliminar grupo Grupo importado',exact=True).click();page.get_by_role('alertdialog').get_by_role('button',name='Eliminar grupo',exact=True).click();expect(page.locator('.editor-groups tbody tr')).to_have_count(4)
 page.get_by_role('button',name='Terminar edición',exact=True).first.click();expect(page.locator('.group-button')).to_have_count(4);page.get_by_text('Personalizada-',exact=False).first.wait_for()
 page.get_by_role('button',name='Consultar existencias',exact=True).click();page.get_by_role('button',name='Seleccionar grupo Grupo nuevo de prueba',exact=True).click();meds.get_by_text('ADRENALINA 1 MG/ML 10 AMPOLLAS',exact=True).wait_for()
 page.get_by_role('button',name='Editar listado',exact=False).click();page.get_by_role('button',name='Restaurar original',exact=True).click();page.get_by_role('alertdialog').get_by_role('button',name='Restaurar original',exact=True).click();expect(page.locator('.editor-groups tbody tr')).to_have_count(3);page.get_by_role('button',name='Terminar edición',exact=True).first.click();expect(page.locator('.group-button')).to_have_count(3);assert fixture.catalog is None
 page.get_by_role('button',name='Consultar existencias',exact=True).click();page.get_by_role('button',name='Seleccionar grupo Grupo ficticio A',exact=True).click();page.get_by_role('button',name='Ver detalle de Grupo ficticio A',exact=True).wait_for()
 page.keyboard.press('Control+k');palette=page.get_by_role('dialog',name='Buscar en Viteka',exact=True);palette.wait_for();page.keyboard.type('650001');palette.get_by_role('option',name=re.compile('en Tenencia')).wait_for();page.screenshot(path=str(artifacts/'angular-palette.png'),full_page=True);page.keyboard.press('Enter');palette.wait_for(state='hidden');expect(page.get_by_label('Buscar grupo, artículo o CN',exact=True)).to_have_value('650001')
 page.keyboard.press('Control+k');palette.wait_for();page.keyboard.type('usuarios');page.keyboard.press('Enter');page.get_by_role('heading',name='Usuarios',exact=True).wait_for()
 # Inicio: información y acceso rápido (sin tarjeta de tenencia); accesos directos en la barra lateral marcados por el usuario.
 page.get_by_role('navigation').get_by_role('button',name='Inicio',exact=True).click();assert page.locator('.tenencia-card').count()==0 and page.get_by_text('correctos',exact=False).count()==0;assert page.get_by_role('navigation').get_by_role('button',name='Tenencia',exact=True).count()==0
 home=page.get_by_role('region',name='Acceso rápido a consultas',exact=True);home.get_by_role('button',name='Favorito',exact=True).click();short=page.get_by_role('navigation').get_by_role('button',name='Acceso directo Tenencia obligatoria',exact=True);short.wait_for();expect(short.locator('.nav-count')).to_have_text('1');page.screenshot(path=str(artifacts/'angular-home-tenencia.png'),full_page=True)
 short.click();page.get_by_role('heading',name='Tenencia obligatoria',exact=True).wait_for();assert short.get_attribute('aria-current')=='page'
 page.get_by_role('navigation').get_by_role('button',name='Inicio',exact=True).click();home.get_by_role('button',name='Favorito',exact=True).click();short.wait_for(state='detached');home.get_by_role('button',name='Abrir Tenencia obligatoria',exact=True).click();page.get_by_role('heading',name='Tenencia obligatoria',exact=True).wait_for()
 go('Aplicaciones');page.get_by_role('button',name='Favorito',exact=True).first.click();go('Favoritos');page.locator('button.resource',has_text='Tenencia obligatoria').wait_for();page.get_by_role('button',name='Favorito',exact=True).first.click();page.get_by_role('heading',name='Aún no tienes favoritos',exact=True).wait_for()
 fixture.use_excel=True;go('Aplicaciones');page.locator('button.resource',has_text='Tenencia obligatoria').click();expect(page.locator('.group-button')).to_have_count(83)
 expect(page.locator('.page-head')).to_have_count(0);expect(page.locator('footer')).to_have_count(0)
 for size in [{'width':1366,'height':768},{'width':1920,'height':1080}]:
  page.set_viewport_size(size)
  page.wait_for_timeout(100)
  geometry=page.evaluate('''() => {const table=document.querySelector('.tenencia-workspace').getBoundingClientRect(),info=document.querySelector('.tenencia-information').getBoundingClientRect();return {height:table.height,top:table.top,bottom:info.bottom,width:document.documentElement.scrollWidth,viewport:innerWidth,viewportHeight:innerHeight,scroll:document.documentElement.scrollHeight};}''')
  assert geometry['height']>=280 and geometry['top']<250 and geometry['bottom']<=size['height'] and geometry['width']<=size['width'] and geometry['scroll']<=size['height']+1,geometry
  page.screenshot(path=str(artifacts/f"angular-tenencia-{size['width']}.png"),full_page=True)
 page.set_viewport_size({'width':1366,'height':768});page.evaluate("document.documentElement.dataset.theme='light'");page.screenshot(path=str(artifacts/'angular-tenencia-light.png'),full_page=True);page.evaluate("document.documentElement.dataset.theme='dark'");page.set_viewport_size({'width':1240,'height':820})
 page.get_by_role('button',name='Seleccionar artículo 152195',exact=True).click();page.get_by_role('region',name='Información de la selección',exact=True).get_by_text('Artículo · CN 152195',exact=True).wait_for();page.keyboard.press('ArrowDown');page.get_by_role('region',name='Información de la selección',exact=True).get_by_text('Artículo · CN 349696',exact=True).wait_for()
 expect(page.locator('.groups-table tbody tr').filter(has=page.get_by_role('button',name='Seleccionar grupo DIGOXINA ORAL 2 ENVASES',exact=True)).locator('.group-code')).to_have_text('C01AA05');page.get_by_role('button',name='Seleccionar grupo DIGOXINA ORAL 2 ENVASES',exact=True).wait_for();page.get_by_role('button',name='Seleccionar grupo DIGOXINA INYECTABLE 1 ENVASE',exact=True).click();page.get_by_role('region',name='Información de la selección',exact=True).get_by_text('C01AA05',exact=True).wait_for()
 page.get_by_role('button',name='Seleccionar grupo HOJA BISTURÍ CON MANGO 1 UNIDAD',exact=True).click();page.get_by_role('region',name='Medicamentos del grupo',exact=True).get_by_text('175550',exact=True).wait_for();expect(page.locator('.tenencia-medications tbody tr')).to_have_count(10);page.get_by_role('button',name='Ver todos (18)',exact=True).click();expect(page.locator('.tenencia-medications tbody tr')).to_have_count(18);page.screenshot(path=str(artifacts/'angular-tenencia-excel.png'),full_page=True);page.get_by_role('button',name='Seleccionar grupo BOLSA DE RECOGIDA DE ORINA PEDIÁTRICA 2 UNIDADES',exact=True).click();expect(page.locator('.tenencia-medications tbody tr')).to_have_count(3);page.get_by_role('region',name='Medicamentos del grupo',exact=True).get_by_text('165727',exact=True).wait_for();fixture.use_excel=False
 go('Usuarios');page.get_by_role('button',name='Crear usuario',exact=True).click();editor=page.get_by_role('dialog',name='Editar cuenta y permisos',exact=True)
 editor.get_by_label('Nombre de usuario',exact=True).fill('Personal');editor.get_by_label('Contraseña inicial',exact=True).fill('FixtureStaffPass1');editor.get_by_label('Repetir contraseña inicial',exact=True).fill('FixtureStaffPass1');editor.get_by_label('Aplicación de prueba A',exact=True).check();editor.get_by_role('button',name='Guardar cuenta',exact=True).click();editor.wait_for(state='hidden');page.get_by_role('button',name='Editar Personal',exact=True).wait_for()
 page.screenshot(path=str(artifacts/'angular-users.png'),full_page=True)
 # Configuración → Utilidades: habilitar utilidades por usuario.
 page.get_by_role('group',name='Secciones de configuración',exact=True).get_by_role('button',name='Utilidades',exact=True).click();page.get_by_role('heading',name='Utilidades por usuario',exact=True).wait_for()
 box=page.get_by_label('Tenencia obligatoria para usuario',exact=True);expect(box).not_to_be_checked();expect(page.get_by_label('Tenencia obligatoria para F. titular',exact=True)).to_be_disabled();box.check();page.get_by_text('Utilidad habilitada para usuario.',exact=True).wait_for();assert 'tenencia-obligatoria' in fixture.accounts[2]['utilityIds'];page.screenshot(path=str(artifacts/'angular-grants.png'),full_page=True)
 box.uncheck();page.get_by_text('Utilidad retirada a usuario.',exact=True).wait_for();assert fixture.accounts[2]['utilityIds']==[]
 go('Mi cuenta');appearance('light')
 page.get_by_role('button',name='Cerrar sesión',exact=True).click();page.get_by_role('heading',name='Iniciar sesión',exact=True).wait_for()
 login('F. titular','1234','FixtureAdminPass1')
 assert page.get_by_role('navigation').get_by_role('button',name='Configuración',exact=True).count()==1
 go('Configuración');assert page.locator('.setting-row').filter(has=page.get_by_text('Conexión a NixFarma',exact=True)).count()==0
 go('Usuarios');assert page.get_by_role('button',name='Editar Viteka',exact=True).count()==0
 page.get_by_role('button',name='Crear usuario',exact=True).click();editor=page.get_by_role('dialog',name='Editar cuenta y permisos',exact=True);editor.wait_for();expect(editor.get_by_label('Rol',exact=True).locator('option')).to_have_count(1);page.keyboard.press('Escape')
 page.get_by_role('button',name='Cerrar sesión',exact=True).click();page.get_by_role('heading',name='Iniciar sesión',exact=True).wait_for()
 login('usuario','0000','FixtureUserPass1')
 assert page.get_by_role('navigation').get_by_role('button',name='Usuarios',exact=True).count()==0
 assert page.get_by_role('navigation').get_by_role('button',name='Configuración',exact=True).count()==1
 go('Configuración');assert page.locator('.setting-row').filter(has=page.get_by_text('Conexión a NixFarma',exact=True)).count()==0
 assert page.locator('.setting-row').filter(has=page.get_by_text('Auditoría',exact=True)).count()==0
 go('Aplicaciones');page.get_by_role('heading',name='No hay aplicaciones disponibles para tu cuenta',exact=True).wait_for()
 fixture.current['utilityIds']=['tenencia-obligatoria'];page.reload();go('Aplicaciones');page.locator('button.resource',has_text='Tenencia obligatoria').click();page.get_by_role('heading',name='Tenencia obligatoria',exact=True).wait_for()
 assert page.get_by_role('button',name='Editar listado',exact=False).count()==0;page.get_by_text(re.compile('Última consulta: 28/10/2026')).wait_for()
 go('Aplicaciones');page.get_by_role('button',name='Favorito',exact=True).click();page.get_by_role('navigation').get_by_role('button',name='Acceso directo Tenencia obligatoria',exact=True).click();page.get_by_role('heading',name='Tenencia obligatoria',exact=True).wait_for()
 go('Actualizaciones');assert page.get_by_role('button',name='Actualizar ahora',exact=True).count()==0;page.get_by_text('El titular o Viteka pueden ejecutar esta actualización desde el servidor de la farmacia.',exact=True).wait_for()
 go('Mi cuenta');page.locator('.setting-row').filter(has=page.get_by_text('Cambiar mi contraseña',exact=True)).click();page.get_by_label('Contraseña actual',exact=True).fill('FixtureUserPass1');page.get_by_label('Nueva contraseña',exact=True).fill('abc');page.get_by_label('Repetir nueva contraseña',exact=True).fill('abc');page.get_by_role('button',name='Cambiar mi contraseña',exact=True).click();page.get_by_text('Contraseña actualizada.',exact=True).wait_for()
 expect(page.get_by_role('dialog')).to_have_count(0)
 page.locator('.setting-row').filter(has=page.get_by_text('Apariencia',exact=True)).click();page.get_by_role('dialog',name='Apariencia',exact=True).get_by_label('Apariencia',exact=True).select_option('dark');page.get_by_role('button',name='Cancelar',exact=True).click();assert page.evaluate('document.documentElement.dataset.theme')=='light'
 page.screenshot(path=str(artifacts/'angular-account-rows.png'),full_page=True)
 page.locator('.setting-row').filter(has=page.get_by_text('Cambiar mi contraseña',exact=True)).click();page.screenshot(path=str(artifacts/'angular-password-dialog.png'),full_page=True);page.keyboard.press('Escape')
 page.set_viewport_size({'width':820,'height':600});assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
 assert not errors,errors
 browser.close()
server.shutdown()
print('PASS: shared pharmacy profile; provincial catalog pending; explicit tenencia query; results, filters and detail; stale results cleared on error; utility access grant; login and first-password change; role navigation; per-user preferences; hidden Oracle configuration; connection messages; users and grants editor; administrator restrictions; update tab and history; escaped release notes; own-password change; compact layout; expiry badges; Excel/CSV/PDF export; top-10 articles; catalog editor; last query; quick access and sidebar shortcuts; utility grants; no JS errors.')
